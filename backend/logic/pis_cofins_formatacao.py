"""
Formatação visual compartilhada pelos módulos de Débitos, Créditos e
Apuração — cabeçalho destacado, painel congelado, largura de coluna
proporcional ao conteúdo do cabeçalho, bordas leves. Feito pra ser
barato mesmo em planilhas com milhares de linhas (não itera célula a
célula em todas as linhas de dado, só no cabeçalho e larguras).
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _formatar_planilha(ws, header_row=1, freeze_cell="A2",
                        header_fill="1F4E78", header_font_color="FFFFFF",
                        largura_extra_col=None, max_col=None):
    """
    ws: worksheet a formatar
    header_row: linha do cabeçalho a destacar
    freeze_cell: célula de referência pro congelamento de painel
    header_fill: cor de fundo do cabeçalho (hex, sem #)
    header_font_color: cor da fonte do cabeçalho (hex, sem #)
    largura_extra_col: tupla opcional (letra_coluna, largura) pra forçar
        uma coluna específica (ex: uma coluna de texto longo)
    max_col: até qual coluna aplicar (padrão: max_column da planilha)
    """
    max_col = max_col or ws.max_column

    fill = PatternFill("solid", fgColor=header_fill)
    font = Font(bold=True, color=header_font_color, size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda_fina = Side(style="thin", color="B7B7B7")
    borda = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

    for c in range(1, max_col + 1):
        cell = ws.cell(header_row, c)
        if cell.value is None:
            continue
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = borda

    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = freeze_cell

    for c in range(1, max_col + 1):
        letra = get_column_letter(c)
        texto = ws.cell(header_row, c).value
        largura = max(10, min(28, len(str(texto)) + 4)) if texto else 10
        ws.column_dimensions[letra].width = largura

    if largura_extra_col:
        letra, largura = largura_extra_col
        ws.column_dimensions[letra].width = largura


def _destacar_totais(ws, linhas, cor_fundo="FFF2CC", negrito=True, max_col=None, min_col=1):
    """Aplica destaque visual (fundo + negrito) numa ou mais linhas de
    total/resumo — não mexe no valor, só na aparência."""
    max_col = max_col or ws.max_column
    fill = PatternFill("solid", fgColor=cor_fundo)
    for r in linhas:
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            if negrito:
                cell.font = Font(bold=True)
            if cell.value is not None:
                cell.fill = fill
