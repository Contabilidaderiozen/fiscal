"""
PROCESSADOR PIS/COFINS - DÉBITOS (versão web, parametrizada)
Preserva a lógica de processar_debitos.py, mas recebe arquivos explícitos
em vez de escanear uma pasta.
"""
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import copy, os, zipfile, re, io
from datetime import datetime
from collections import OrderedDict
import xlrd as _xlrd

from . import cfop_classificacao
CFOP_MAP = cfop_classificacao.CFOP_MAP
from .pis_cofins_formatacao import _formatar_planilha, _destacar_totais

FMT_MOEDA = '#,##0.00_);[Red](#,##0.00)'
FMT_PCT   = '0.00%'


def limpar_ncm(val):
    if pd.isna(val) or str(val).strip() == '': return ''
    try: return str(int(float(str(val).strip())))
    except: return str(val).strip()


def limpar_conta(val):
    return str(val).replace('.', '').strip() if val else val


def safe(val):
    return None if pd.isna(val) else val


CHAVES_RESUMO = {
    ('Peças e Acessórios', 'Normal NC'):    'PEÇAS E ACESSÓRIOS - Normal NC',
    ('Garantia',           'Normal NC'):    'PEÇAS E ACESSÓRIOS - Normal NC',
    ('Veículo Usado',      'Normal C'):     'VEICULO USADO - NORMAL C',
    ('Peças e Acessórios', 'Monofasico'):   'PEÇAS E ACESSÓRIOS - Monofasico',
    ('Veículo Novo',       'Monofasico'):   'VEICULO NOVO - MONOFASICO',
    ('Veículo Novo',       'Zero'):         'VEICULO NOVO - MONOFASICO',
    ('Garantia',           'Zero'):         'GARANTIA - ZERO',
    ('Ativo Imobilizado',  'Zero'):         'ATIVO IMOBILIZADO - ZERO',
    ('Devolução',          'Zero'):         'DEVOLUÇÃO - ZERO',
    ('Comissão',    'Normal NC'): 'COMISSÃO - NORMAL NC',
    ('Oficina',     'Normal NC'): 'OFICINA - NORMAL NC',
    ('Venda Direta','Zero'):      'VENDA DIRETA - ZERO',
    ('Combustível', 'Normal NC'): 'COMBUSTÍVEL - Normal NC',
    ('Combustível', 'Monofasico'): 'COMBUSTÍVEL - Monofasico',
    ('Material de Uso e Consumo', 'Normal NC'): 'MATERIAL USO E CONSUMO - Normal NC',
    ('Material de Uso e Consumo', 'Monofasico'): 'MATERIAL USO E CONSUMO - Monofasico',
    ('Serviço - A Classificar', 'Normal NC'): 'A CLASSIFICAR - NORMAL NC',
}


def classificar_operacao(cfop, ncm):
    cfop_str = str(int(float(cfop))) if pd.notna(cfop) and str(cfop) != 'nan' else ''
    tem_ncm  = pd.notna(ncm) and str(ncm).strip() not in ('', 'nan', '0')
    if cfop_str == '':
        return 'Serviço', 'Serviço - A Classificar'
    if cfop_str in ('5102', '6102'):
        return 'Venda', ('Peças e Acessórios' if tem_ncm else 'Veículo Usado')
    if cfop_str == '5405':
        return 'Venda', ('Peças e Acessórios' if tem_ncm else 'Veículo Novo')
    entry = CFOP_MAP.get(cfop_str)
    if entry:
        return entry[0], entry[1]
    return '', ''


def chave_resumo(tipo_h, tipo_y):
    return CHAVES_RESUMO.get((tipo_h, tipo_y), f'{tipo_h.upper()} - {tipo_y.upper()}')


def ler_arquivo_170(arq170, log):
    mapa_nota_tipo = {}
    if not arq170:
        log("[0] Arquivo 170 não fornecido — classificação por CST Pis")
        return mapa_nota_tipo
    log(f"[0] Lendo arquivo 170: {os.path.basename(arq170)}")
    try:
        ext170 = os.path.splitext(arq170)[1].lower()
        if ext170 in ('.xlsx', '.xls'):
            engine = 'openpyxl' if ext170 == '.xlsx' else 'xlrd'
            # O arquivo 170 pode vir com várias abas (uma por tipo de
            # registro SPED: 0000, A100, A170...) ou com tudo numa aba só
            # (formato antigo). Os dados de item (nota + descrição) sempre
            # ficam no registro A170, então procuramos essa aba pelo nome
            # antes de cair pra primeira aba como fallback.
            _xf = pd.ExcelFile(arq170, engine=engine)
            _aba_alvo = next((s for s in _xf.sheet_names if s.strip().upper() == 'A170'), _xf.sheet_names[0])
            df170 = pd.read_excel(_xf, sheet_name=_aba_alvo, header=None, dtype=str)
        else:
            df170 = pd.read_csv(arq170, sep=';', encoding='latin1', header=None, dtype=str, skiprows=1)
        df170.columns = range(len(df170.columns))
        IDX_NOTA, IDX_DESC = 14, 31
        for _, row170 in df170.iterrows():
            nf = ''
            if IDX_NOTA < len(row170):
                nf = str(row170.iloc[IDX_NOTA]).strip().lstrip("'").strip()
            if not nf or nf.upper() in ('NAN', 'NONE', '', 'NUM_DOC-A100'):
                continue
            desc = ''
            if IDX_DESC < len(row170):
                desc = str(row170.iloc[IDX_DESC]).strip().lstrip("'").strip().upper()
            if 'VENDA DIRETA' in desc:
                tipo_serv = 'Venda Direta'
            elif 'COMISS' in desc and 'VENDA' in desc:
                tipo_serv = 'Venda Direta'
            elif 'COMISS' in desc:
                tipo_serv = 'Comissão'
            else:
                continue
            atual = mapa_nota_tipo.get(nf)
            if atual != 'Venda Direta':
                mapa_nota_tipo[nf] = tipo_serv
        log(f"   {len(mapa_nota_tipo)} notas mapeadas: "
            f"Venda Direta={sum(1 for v in mapa_nota_tipo.values() if v=='Venda Direta')}, "
            f"Comissão={sum(1 for v in mapa_nota_tipo.values() if v=='Comissão')}")
    except Exception as e:
        log(f"   AVISO: erro ao ler 170: {e}")
    return mapa_nota_tipo


def processar_debitos(xls_files, modelo, balancete, arquivo170, saida, log=print):
    mapa_nota_tipo = ler_arquivo_170(arquivo170, log)

    log("\n" + "="*60)
    log("PROCESSADOR PIS/COFINS - DÉBITOS")
    log("="*60)

    log("\n[1] Lendo arquivos XLS...")
    if not xls_files:
        raise ValueError("Nenhum arquivo de filial fornecido!")

    all_dfs = []
    for path in xls_files:
        fname = os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()
        engine = "xlrd" if ext == ".xls" else "openpyxl"
        df = pd.read_excel(path, engine=engine, skiprows=5, header=0, sheet_name=0)
        df = df.dropna(how='all').reset_index(drop=True)
        while len(df.columns) < 19:
            df[len(df.columns)] = None
        for ci in [0, 1, 2]:
            df.iloc[:, ci] = df.iloc[:, ci].ffill()

        col_r = df.columns[16]
        col_q = df.columns[17]
        col_s = df.columns[18]
        mask = df[col_r].notna() & (df[col_r] != 0)
        df.loc[mask, col_s] = df.loc[mask, col_q]
        df.loc[mask, col_q] = df.loc[mask, col_r]
        df.loc[mask, col_r] = None
        log(f"   Corrigidos {mask.sum()} linhas de Base/Valor ICMS")

        ops, tipos = [], []
        for _, row in df.iterrows():
            cfop = row.iloc[5]
            ncm  = row.iloc[4]
            cod  = str(row.iloc[3]) if pd.notna(row.iloc[3]) else ''
            if cod.strip().upper() == 'SERV':
                nf_str = str(row.iloc[0]).strip().lstrip("'").split('.')[0].strip()
                tipo_170 = mapa_nota_tipo.get(nf_str)
                if tipo_170 == 'Venda Direta':
                    op, tipo = 'Venda', 'Venda Direta'
                else:
                    cst_pis = row.iloc[6] if len(row) > 6 else None
                    cst_val = None
                    if pd.notna(cst_pis) and str(cst_pis).strip() not in ('', 'nan', '0.0'):
                        try: cst_val = int(float(str(cst_pis)))
                        except Exception: cst_val = None
                    if cst_val == 1:
                        op, tipo = 'Serviço', 'Oficina'
                    else:
                        op, tipo = 'Serviço', 'Comissão'
            else:
                op, tipo = classificar_operacao(cfop, ncm)
            ops.append(op)
            tipos.append(tipo)
        df.insert(6, 'Operação', ops)
        df.insert(7, 'Tipo Saida', tipos)
        all_dfs.append(df)
        log(f"   OK: {fname} ({len(df)} linhas)")

    if not all_dfs:
        raise ValueError("Nenhum arquivo lido!")

    combined = pd.concat(all_dfs, ignore_index=True)
    combined = combined[combined.iloc[:, 0].notna()].copy().reset_index(drop=True)
    ncm_col = combined.columns[4]
    combined[ncm_col] = combined[ncm_col].apply(lambda x: limpar_ncm(x) if pd.notna(x) else '')
    log(f"   Total de itens: {len(combined)}")

    log("\n[2] Carregando modelo...")
    if not modelo:
        raise ValueError("Modelo não fornecido!")

    ext_modelo = os.path.splitext(modelo)[1].lower()
    if ext_modelo == '.xls':
        rb = _xlrd.open_workbook(modelo)
        buf = io.BytesIO()
        from openpyxl import Workbook as WB2
        wb_tmp = WB2()
        for sh_idx in range(rb.nsheets):
            sh = rb.sheet_by_index(sh_idx)
            ws_tmp = wb_tmp.active if sh_idx == 0 else wb_tmp.create_sheet()
            ws_tmp.title = sh.name
            for rx in range(sh.nrows):
                ws_tmp.append([sh.cell_value(rx, cx) for cx in range(sh.ncols)])
        wb_tmp.save(buf)
        buf.seek(0)
        wb_orig = load_workbook(buf)
    else:
        wb_orig = load_workbook(modelo)
    ws_orig = wb_orig.active

    col_widths = {col: ws_orig.column_dimensions[col].width for col in ws_orig.column_dimensions}
    wb_orig.close()

    # Salvar workbook convertido como ponto de partida
    if ext_modelo == '.xls':
        rb = _xlrd.open_workbook(modelo)
        from openpyxl import Workbook as WB2
        wb = WB2()
        for sh_idx in range(rb.nsheets):
            sh = rb.sheet_by_index(sh_idx)
            ws_tmp = wb.active if sh_idx == 0 else wb.create_sheet()
            ws_tmp.title = sh.name
            for rx in range(sh.nrows):
                ws_tmp.append([sh.cell_value(rx, cx) for cx in range(sh.ncols)])
    else:
        wb = load_workbook(modelo)

    for nome in list(wb.sheetnames):
        if nome.upper() not in ("NCM", "BALANCETE"):
            del wb[nome]
    ws_deb = wb.create_sheet("DEBITOS", 0)
    if "NCM" not in wb.sheetnames:
        wb.create_sheet("NCM")
    ws_ncm = wb["NCM"]

    for col, width in col_widths.items():
        ws_deb.column_dimensions[col].width = width

    cab = [
        'Nota Fiscal','Serie','Emissão','Código do item','NCM','CFOP',
        'Operação','Tipo Saida','Descrição Tipo (I)',
        'CST Pis','CST Cofins','Base Pis','Aliq. Pis','Valor Pis',
        'Base Cofins','Aliq. Cofins','Valor Cofins','Vlr Pis + Vlr Cofins',
        'Fat. Bruto','Base Icms','Valor Icms','Aliq. Efetiva ICMS','',
        'NCM','TIPO','PIS','COFINS','',
        'Valor Icms','CHECK','ICMS 2%','Total ICMS','CHECK'
    ]
    for i, label in enumerate(cab, 1):
        ws_deb.cell(row=1, column=i, value=label).font = Font(bold=True)

    log("\n[3] Atualizando Balancete...")
    bal_max = 0
    bal_nome = next((n for n in wb.sheetnames if n.upper() == "BALANCETE"), None)
    if bal_nome is None:
        wb.create_sheet("BALANCETE")
        bal_nome = "BALANCETE"
    ws_bal = wb[bal_nome]
    bal_max = ws_bal.max_row

    if balancete:
        log(f"   {os.path.basename(balancete)}")
        wb_bal_src = None
        try:
            wb_bal_src = load_workbook(balancete, read_only=True, data_only=True)
        except Exception:
            try:
                with zipfile.ZipFile(balancete, 'r') as zin:
                    xml = zin.read('xl/styles.xml').decode('utf-8')
                counter = [0]
                def fix_style(m):
                    tag = m.group(0)
                    if 'name=' not in tag:
                        tag = tag.replace('<cellStyle ', f'<cellStyle name="B{counter[0]}" ')
                        counter[0] += 1
                    return tag
                xml = re.sub(r'<cellStyle [^/]*/>', fix_style, xml)
                buf = io.BytesIO()
                with zipfile.ZipFile(balancete, 'r') as zin:
                    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
                        for item in zin.infolist():
                            data = xml.encode('utf-8') if item.filename == 'xl/styles.xml' else zin.read(item.filename)
                            zout.writestr(item, data)
                buf.seek(0)
                wb_bal_src = load_workbook(buf, read_only=True, data_only=True)
                log("   (estilos reparados automaticamente)")
            except Exception as e2:
                log(f"   AVISO: nao foi possivel abrir balancete ({e2})")

        if wb_bal_src:
            ws_bs = wb_bal_src.active
            COL_MAP = {1: 1, 2: 2, 3: 4, 4: 7, 5: 8, 6: 9}
            for row in ws_bal.iter_rows(min_row=4, max_row=ws_bal.max_row):
                for cell in row:
                    cell.value = None
            dest = 4
            for i, row in enumerate(ws_bs.iter_rows(min_row=4, values_only=True), start=4):
                if i == 4:
                    headers = ["Conta","Nome da Conta","Saldo Anterior","Devedor","Credor","Movimento do Mês"]
                    for j, h in enumerate(headers, 1):
                        ws_bal.cell(row=dest, column=j, value=h)
                    dest += 1
                    continue
                if row[0] is None: continue
                conta = limpar_conta(row[0])
                try: conta = int(conta)
                except: pass
                for dest_col, orig_col in COL_MAP.items():
                    v = row[orig_col - 1] if orig_col - 1 < len(row) else None
                    ws_bal.cell(row=dest, column=dest_col, value=conta if dest_col == 1 else v)
                dest += 1
            bal_max = dest - 1
            log(f"   {dest-5} linhas (max_row={bal_max})")
    else:
        log("   Balancete nao encontrado - pulando")

    log("\n[4] Verificando NCMs...")
    existing_ncms = {}
    for row in ws_ncm.iter_rows(min_row=2, values_only=True):
        if row[0]: existing_ncms.setdefault(limpar_ncm(row[0]), row)  # 1ª ocorrência vence (igual VLOOKUP)

    missing_ncms = {}
    for _, row in combined.iterrows():
        ncm = limpar_ncm(row.iloc[4])
        if ncm and ncm not in existing_ncms and ncm not in missing_ncms:
            pis = row.iloc[9] if pd.notna(row.iloc[9]) else 0
            cof = row.iloc[12] if pd.notna(row.iloc[12]) else 0
            try: pv, cv = float(pis), float(cof)
            except: pv = cv = 0
            tipo = "Normal NC" if (pv > 0 or cv > 0) else "Monofasico"
            missing_ncms[ncm] = (tipo, pv, cv)

    if missing_ncms:
        nr = ws_ncm.max_row + 1
        for ncm, (tipo, pv, cv) in missing_ncms.items():
            try: ncm_val = int(ncm)
            except: ncm_val = ncm
            ws_ncm.cell(row=nr, column=1, value=ncm_val)
            ws_ncm.cell(row=nr, column=2, value=tipo)
            ws_ncm.cell(row=nr, column=3, value=pv)
            ws_ncm.cell(row=nr, column=4, value=cv)
            nr += 1
        log(f"   {len(missing_ncms)} NCMs adicionados")
    else:
        log("   Todos cadastrados!")

    log("\n[5] Escrevendo dados...")
    for r_idx, (_, row) in enumerate(combined.iterrows()):
        r = r_idx + 2
        ws_deb.cell(r, 1,  value=safe(row.iloc[0]))
        ws_deb.cell(r, 2,  value=safe(row.iloc[1]))
        ws_deb.cell(r, 3,  value=safe(row.iloc[2]))
        ws_deb.cell(r, 4,  value=safe(row.iloc[3]))
        ncm_str = limpar_ncm(row.iloc[4])
        try: ws_deb.cell(r, 5, value=int(ncm_str))
        except: ws_deb.cell(r, 5, value=ncm_str if ncm_str else None)
        ws_deb.cell(r, 6,  value=safe(row.iloc[5]))
        ws_deb.cell(r, 7,  value=safe(row.iloc[6]))
        tipo_h = safe(row.iloc[7]) or ''
        ws_deb.cell(r, 8,  value=tipo_h)
        ws_deb.cell(r, 9,  value=f'__CHAVE_{r}__')
        ws_deb.cell(r, 10, value=safe(row.iloc[8]))
        ws_deb.cell(r, 11, value=safe(row.iloc[9]))
        for col, idx in [(12,10),(13,11),(14,12),(15,13),(16,14),(17,15),(18,16),(19,17),(20,19),(21,20)]:
            cell = ws_deb.cell(r, col, value=safe(row.iloc[idx]))
            cell.number_format = FMT_MOEDA

        tipo_h = safe(row.iloc[7])
        if tipo_h in ('Garantia', 'Ativo Imobilizado', 'Devolução'):
            tipo_y = 'Zero'
            ws_deb.cell(r, 24, value=tipo_h.upper())
            ws_deb.cell(r, 25, value=tipo_y)
            ws_deb.cell(r, 26, value=0)
            ws_deb.cell(r, 27, value=0)
        elif tipo_h == 'Venda Direta':
            tipo_y = 'Zero'
            ws_deb.cell(r, 24, value='VENDA DIRETA')
            ws_deb.cell(r, 25, value=tipo_y)
            ws_deb.cell(r, 26, value=0)
            ws_deb.cell(r, 27, value=0)
        elif tipo_h in ('Comissão', 'Oficina', 'Serviço - A Classificar', 'Combustível', 'Material de Uso e Consumo'):
            tipo_y = 'Normal NC'
            label = tipo_h.upper() if tipo_h != 'Serviço - A Classificar' else 'A CLASSIFICAR'
            ws_deb.cell(r, 24, value=label)
            ws_deb.cell(r, 25, value=tipo_y)
            ws_deb.cell(r, 26, value=0.0165)
            ws_deb.cell(r, 27, value=0.076)
            ws_deb.cell(r, 26).number_format = '0.00%'
            ws_deb.cell(r, 27).number_format = '0.00%'
        elif tipo_h == 'Veículo Novo':
            tipo_y = 'Monofasico'
            ws_deb.cell(r, 24, value='VEICULO NOVO')
            ws_deb.cell(r, 25, value=tipo_y)
            ws_deb.cell(r, 26, value=0)
            ws_deb.cell(r, 27, value=0)
        elif tipo_h == 'Veículo Usado':
            tipo_y = 'Normal C'
            ws_deb.cell(r, 24, value='VEICULO USADO')
            ws_deb.cell(r, 25, value=tipo_y)
            ws_deb.cell(r, 26, value=0.0065)
            ws_deb.cell(r, 27, value=0.03)
            ws_deb.cell(r, 26).number_format = '0.00%'
            ws_deb.cell(r, 27).number_format = '0.00%'
        else:
            ncm_str2 = limpar_ncm(row.iloc[4])
            if ncm_str2 in existing_ncms:
                tipo_y = existing_ncms[ncm_str2][1] or 'Normal NC'
            elif ncm_str2 in missing_ncms:
                tipo_y = missing_ncms[ncm_str2][0]
            else:
                tipo_y = 'Normal NC'
            ws_deb.cell(r, 24, value=f'=IFERROR(VLOOKUP(E{r},NCM!$A:$D,1,0),"")')
            ws_deb.cell(r, 25, value=f'=IFERROR(VLOOKUP(E{r},NCM!$A:$D,2,0),"")')
            ws_deb.cell(r, 26, value=f'=IFERROR(VLOOKUP(E{r},NCM!$A:$D,3,0),"")')
            ws_deb.cell(r, 27, value=f'=IFERROR(VLOOKUP(E{r},NCM!$A:$D,4,0),"")')
            ws_deb.cell(r, 20, value=0)
            ws_deb.cell(r, 21, value=0)
        ws_deb.cell(r, 22, value=f'=IFERROR(U{r}/T{r},0)').number_format = '0.00%'
        ws_deb.cell(r, 29, value=f'=U{r}').number_format = FMT_MOEDA
        ws_deb.cell(r, 30, value=f'=IFERROR(AC{r}-(T{r}*V{r}),0)').number_format = FMT_MOEDA
        ws_deb.cell(r, 31, value=f'=IFERROR(T{r}*0.02,0)').number_format = FMT_MOEDA
        ws_deb.cell(r, 32, value=f'=IFERROR(AC{r}+AE{r},0)').number_format = FMT_MOEDA
        ws_deb.cell(r, 33, value=f'=IFERROR(AF{r},0)').number_format = FMT_MOEDA
        ws_deb.cell(r, 9, value=chave_resumo(tipo_h, tipo_y))

    last_data_row = len(combined) + 1
    log(f"   Itens escritos: {len(combined)} | Ultima linha: {last_data_row}")

    log("\n[6] Gerando resumo...")
    chaves_normal_nc  = OrderedDict()
    chaves_cumulativo = OrderedDict()
    chaves_zero       = OrderedDict()
    for _, row in combined.iterrows():
        tipo_h = str(row.iloc[7]) if pd.notna(row.iloc[7]) else ''
        if tipo_h in ('Garantia', 'Ativo Imobilizado', 'Devolução', 'Venda Direta'):
            tipo_y = 'Zero'
        elif tipo_h in ('Comissão', 'Oficina', 'Serviço - A Classificar', 'Combustível', 'Material de Uso e Consumo'):
            tipo_y = 'Normal NC'
        elif tipo_h == 'Veículo Novo':
            tipo_y = 'Monofasico'
        elif tipo_h == 'Veículo Usado':
            tipo_y = 'Normal C'
        else:
            ncm_str2 = limpar_ncm(row.iloc[4])
            if ncm_str2 in existing_ncms:
                tipo_y = existing_ncms[ncm_str2][1] or 'Normal NC'
            elif ncm_str2 in missing_ncms:
                tipo_y = missing_ncms[ncm_str2][0]
            else:
                tipo_y = 'Normal NC'
        chave = chave_resumo(tipo_h, tipo_y)
        if not chave: continue
        if tipo_y in ('Zero', 'Monofasico'):
            chaves_zero[chave] = True
        elif tipo_y == 'Normal C':
            chaves_cumulativo[chave] = True
        else:
            chaves_normal_nc[chave] = True

    normal_nc  = sorted(chaves_normal_nc.keys())
    cumulativo = sorted(chaves_cumulativo.keys())
    zero_mono  = sorted(chaves_zero.keys())
    log(f"   Normal NC:  {normal_nc}")
    log(f"   Cumulativo: {cumulativo}")
    log(f"   Zero/Mono:  {zero_mono}")

    KEY_COL = 'I'
    resumo_row   = last_data_row + 3
    subtotal_row = resumo_row + 1

    headers_res = [
        (12,"Base Pis"),(13,"Alíquota Pis"),(14,"Valor Pis"),
        (15,"Base Cofins"),(16,"Alíquota Cofins"),(17,"Valor Cofins"),
        (18,"Total Recolher"),(19,"Fat. Bruto"),
        (20,"Base ICMS"),(21,"Valor Icms"),(22,"2%"),(23,"Icms total")
    ]
    for col, label in headers_res:
        ws_deb.cell(row=resumo_row, column=col, value=label).font = Font(bold=True)

    for col, rng in [(12,"L"),(14,"N"),(15,"O"),(17,"Q"),(18,"R"),(19,"S"),(20,"T"),(21,"U")]:
        ws_deb.cell(row=subtotal_row, column=col,
                    value=f"=SUBTOTAL(9,{rng}2:{rng}{last_data_row})").number_format = FMT_MOEDA
    ws_deb.cell(row=subtotal_row, column=22, value=f"=T{subtotal_row}*0.02").number_format = FMT_MOEDA
    ws_deb.cell(row=subtotal_row, column=23, value=f"=U{subtotal_row}+V{subtotal_row}").number_format = FMT_MOEDA

    def escrever_tabela_deb(header_row, descricoes, aliq_pis, aliq_cof):
        ws_deb.cell(row=header_row, column=11, value="Descrição").font = Font(bold=True)
        for col, label in headers_res:
            ws_deb.cell(row=header_row, column=col, value=label).font = Font(bold=True)
        first = header_row + 1
        for i, desc in enumerate(descricoes):
            r = first + i
            ws_deb.cell(r, 11, value=desc)
            ws_deb.cell(r, 12, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},L$2:L${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 13, value=aliq_pis).number_format = FMT_PCT
            ws_deb.cell(r, 14, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},N$2:N${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 15, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},O$2:O${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 16, value=aliq_cof).number_format = FMT_PCT
            ws_deb.cell(r, 17, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},Q$2:Q${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 18, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},R$2:R${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 19, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},S$2:S${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 20, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},T$2:T${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 21, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},U$2:U${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 22, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},AE$2:AE${last_data_row}),0)').number_format = FMT_MOEDA
            ws_deb.cell(r, 23, value=f'=IFERROR(SUMIF(${KEY_COL}$2:${KEY_COL}${last_data_row},K{r},AG$2:AG${last_data_row}),0)').number_format = FMT_MOEDA
        last  = first + len(descricoes) - 1
        total = last + 1
        fill_amarelo = PatternFill("solid", fgColor="FFFF00")
        for col in range(11, 24):
            ws_deb.cell(total, col).fill = fill_amarelo
            ws_deb.cell(total, col).font = Font(bold=True)
        ws_deb.cell(total, 11, value="Total")
        for col, rng in [(12,'L'),(14,'N'),(15,'O'),(17,'Q'),(18,'R'),(19,'S'),(20,'T'),(21,'U'),(22,'V'),(23,'W')]:
            ws_deb.cell(total, col, value=f'=SUM({rng}{first}:{rng}{last})').number_format = FMT_MOEDA
        for col in range(11, 24):
            ws_deb.cell(total, col).fill = fill_amarelo
            ws_deb.cell(total, col).font = Font(bold=True)
        for col in [13, 16]:
            ws_deb.cell(total, col, value='-')
        return total

    if not normal_nc: normal_nc = ['(sem lançamentos)']
    if not cumulativo: cumulativo = ['(sem lançamentos)']
    if not zero_mono: zero_mono = ['(sem lançamentos)']

    nc_header  = resumo_row + 3
    nc_total   = escrever_tabela_deb(nc_header,  normal_nc,  0.0165, 0.076)
    cum_header = nc_total + 2
    cum_total  = escrever_tabela_deb(cum_header, cumulativo, 0.0065, 0.03)
    zm_header  = cum_total + 2
    zm_total   = escrever_tabela_deb(zm_header,  zero_mono,  0,      0)

    tg_row    = zm_total + 2
    check_row = tg_row + 1
    ws_deb.cell(tg_row,    11, value="Total Geral").font = Font(bold=True)
    ws_deb.cell(check_row, 11, value="Check").font = Font(bold=True)
    for col, c in [(12,'L'),(14,'N'),(15,'O'),(17,'Q'),(18,'R'),(19,'S'),(20,'T'),(21,'U'),(22,'V'),(23,'W')]:
        ws_deb.cell(tg_row,    col, value=f'={c}{nc_total}+{c}{cum_total}+{c}{zm_total}').number_format = FMT_MOEDA
        ws_deb.cell(check_row, col, value=f'={c}{tg_row}-{c}{subtotal_row}').number_format = FMT_MOEDA
    ws_deb.cell(resumo_row, 22, value="2%").font = Font(bold=True)
    ws_deb.cell(resumo_row, 23, value="Icms total").font = Font(bold=True)
    for col in [13, 16]:
        ws_deb.cell(tg_row,    col, value='-')
        ws_deb.cell(check_row, col, value='-')

    log("\n[7] Gerando seção fixa com VLOOKUP no Balancete...")
    ext_modelo2 = os.path.splitext(modelo)[1].lower()
    if ext_modelo2 == '.xls':
        _wb_orig_xls = _xlrd.open_workbook(modelo)
        _sh = _wb_orig_xls.sheet_by_index(0)
        _nrows, _ncols = _sh.nrows, _sh.ncols
        def _cell_val(r, c): return _sh.cell_value(r, c)
    else:
        _wb_orig_xlsx = load_workbook(modelo, data_only=True)
        _sh = _wb_orig_xlsx.active
        _nrows, _ncols = _sh.max_row, _sh.max_column
        def _cell_val(r, c):
            v = _sh.cell(r+1, c+1).value
            return v if v is not None else ''

    secao_inicio_xls = None
    for _r in range(_nrows):
        if _cell_val(_r, 10) == 'DESCRIÇÃO' and 'BASE' in str(_cell_val(_r, 12)).upper():
            secao_inicio_xls = _r
            break

    if secao_inicio_xls is None:
        log("   AVISO: seção fixa não encontrada no modelo")
    else:
        dest_secao = check_row + 3
        bal_ref = f"BALANCETE!$A$5:$F${bal_max}"
        r_dest = dest_secao
        for _r in range(secao_inicio_xls, _nrows):
            row_vals = [_cell_val(_r, c) for c in range(_ncols)]
            non_empty = [(c, v) for c, v in enumerate(row_vals) if v != '']
            if not non_empty:
                r_dest += 1
                continue

            desc  = row_vals[10]
            conta = row_vals[12]
            aliq_pis  = row_vals[14] if row_vals[14] != '' else None
            aliq_cof  = row_vals[17] if row_vals[17] != '' else None

            if desc in ('DESCRIÇÃO',) or str(conta) in ('CONTA', 'BASE SEM EMISSÃO DE NOTA FISCAL', 'BASE RECEITA FINANCEIRA'):
                ws_deb.cell(r_dest, 11, value=desc).font = Font(bold=True)
                ws_deb.cell(r_dest, 13, value=str(conta) if conta else '').font = Font(bold=True)
                for _c, _label in [(14,'Base Pis'),(15,'Aliq. Pis'),(16,'Valor Pis'),
                                   (17,'Base Cofins'),(18,'Aliq. Cofins'),(19,'Valor Cofins'),
                                   (20,'Fat. Bruto'),(21,'Base Icms')]:
                    ws_deb.cell(r_dest, _c, value=_label).font = Font(bold=True)
                r_dest += 1
                continue

            if str(desc).startswith('TOTAL') or str(conta).startswith('TOTAL'):
                lbl = str(desc) if str(desc).startswith('TOTAL') else str(conta)
                if 'OUTRAS RECEITAS' in lbl.upper() and not lbl.upper().startswith('TOTAL OUTRAS RECEITAS 35'):
                    r_dest += 1
                    continue
                grp_end = r_dest - 1
                grp_start = grp_end
                while grp_start > dest_secao:
                    v = ws_deb.cell(grp_start-1, 13).value
                    if v is None or str(v) in ('CONTA', '') or str(v).upper().startswith('BASE') or str(v).upper().startswith('TOTAL'):
                        break
                    grp_start -= 1
                fill_c = PatternFill("solid", fgColor="D9D9D9")
                if lbl.upper().startswith('TOTAL GERAL'):
                    totais_rows = []
                    for _tr in range(dest_secao, r_dest):
                        _v = ws_deb.cell(_tr, 13).value
                        if _v and (str(_v).upper().startswith('TOTAL OUTRAS RECEITAS') or
                                   str(_v).upper().startswith('TOTAL RECEITA')):
                            totais_rows.append(_tr)
                    ws_deb.cell(r_dest, 13, value='TOTAL GERAL').font = Font(bold=True)
                    for _c in range(11, 25):
                        ws_deb.cell(r_dest, _c).fill = fill_c
                        ws_deb.cell(r_dest, _c).font = Font(bold=True)
                    if len(totais_rows) >= 2:
                        t1, t2 = totais_rows[0], totais_rows[1]
                        for _c, _rng in [(14,'N'),(16,'P'),(17,'Q'),(19,'S'),(20,'T')]:
                            ws_deb.cell(r_dest, _c, value=f'={_rng}{t1}+{_rng}{t2}').number_format = FMT_MOEDA
                    elif len(totais_rows) == 1:
                        t1 = totais_rows[0]
                        for _c, _rng in [(14,'N'),(16,'P'),(17,'Q'),(19,'S'),(20,'T')]:
                            ws_deb.cell(r_dest, _c, value=f'={_rng}{t1}').number_format = FMT_MOEDA
                    ws_deb.cell(r_dest, 15, value='-')
                    ws_deb.cell(r_dest, 18, value='-')
                    r_dest += 1
                    continue
                if lbl.upper() == 'CHECK':
                    ws_deb.cell(r_dest, 13, value='CHECK').font = Font(bold=True)
                    for _c in range(11, 25):
                        ws_deb.cell(r_dest, _c).fill = fill_c
                        ws_deb.cell(r_dest, _c).font = Font(bold=True)
                    r_dest += 1
                    continue
                ws_deb.cell(r_dest, 13, value=lbl).font = Font(bold=True)
                for _c in range(11, 25):
                    ws_deb.cell(r_dest, _c).fill = fill_c
                    ws_deb.cell(r_dest, _c).font = Font(bold=True)
                if grp_start <= grp_end:
                    for _c, _rng in [(14,'N'),(16,'P'),(17,'Q'),(19,'S'),(20,'T')]:
                        ws_deb.cell(r_dest, _c, value=f'=SUM({_rng}{grp_start}:{_rng}{grp_end})').number_format = FMT_MOEDA
                    ws_deb.cell(r_dest, 15, value='-')
                    ws_deb.cell(r_dest, 18, value='-')
                r_dest += 1
                r_dest += 1
                continue

            if conta and str(conta) not in ('', 'nan') and str(desc) not in ('', 'nan'):
                try: conta_int = int(float(str(conta)))
                except: conta_int = conta
                ws_deb.cell(r_dest, 11, value=str(desc))
                ws_deb.cell(r_dest, 13, value=conta_int)
                vlookup = f'=IFERROR(VLOOKUP(M{r_dest},{bal_ref},6,0),0)'
                ws_deb.cell(r_dest, 14, value=vlookup).number_format = FMT_MOEDA
                ws_deb.cell(r_dest, 15, value=aliq_pis if aliq_pis else 0).number_format = FMT_PCT
                ws_deb.cell(r_dest, 16, value=f'=N{r_dest}*O{r_dest}').number_format = FMT_MOEDA
                ws_deb.cell(r_dest, 17, value=f'=N{r_dest}').number_format = FMT_MOEDA
                ws_deb.cell(r_dest, 18, value=aliq_cof if aliq_cof else 0).number_format = FMT_PCT
                ws_deb.cell(r_dest, 19, value=f'=R{r_dest}*Q{r_dest}').number_format = FMT_MOEDA
                ws_deb.cell(r_dest, 20, value=f'=N{r_dest}').number_format = FMT_MOEDA
                ws_deb.cell(r_dest, 21, value=0).number_format = FMT_MOEDA
                r_dest += 1
        log(f"   Seção fixa gerada: linhas {dest_secao} a {r_dest-1}")

    ws_deb.auto_filter.ref = f"A1:AA{last_data_row}"

    prop_row = check_row + 2
    for col, label in [(24, 'Descrição'), (25, 'Valor'), (26, '%')]:
        cell = ws_deb.cell(prop_row, col, value=label)
        cell.font = Font(bold=True)
    fat_row    = prop_row + 1
    usado_row  = prop_row + 2
    demais_row = prop_row + 3
    ws_deb.cell(fat_row, 24, value='FAT BRUTO').font = Font(bold=True)
    ws_deb.cell(fat_row, 25, value=f'=S{subtotal_row}').number_format = FMT_MOEDA
    ws_deb.cell(fat_row, 26, value=1.0).number_format = '0%'
    ws_deb.cell(usado_row, 24, value='VEIC USADO').font = Font(bold=True)
    ws_deb.cell(usado_row, 25,
        value=f'=SUMIF($I$2:$I${last_data_row},"VEICULO USADO - NORMAL C",$S$2:$S${last_data_row})').number_format = FMT_MOEDA
    ws_deb.cell(usado_row, 26,
        value=f'=IF(Y{fat_row}=0,0,Y{usado_row}/Y{fat_row})').number_format = '0%'
    ws_deb.cell(demais_row, 24, value='DEMAIS').font = Font(bold=True)
    ws_deb.cell(demais_row, 25, value=f'=Y{fat_row}-Y{usado_row}').number_format = FMT_MOEDA
    ws_deb.cell(demais_row, 26,
        value=f'=IF(Y{fat_row}=0,0,Y{demais_row}/Y{fat_row})').number_format = '0%'
    log(f"   Tabela proporcional escrita nas linhas {fat_row}-{demais_row} (cols X,Y,Z)")

    log("\n[8] Aplicando formatação visual...")
    _formatar_planilha(ws_deb, header_row=1, freeze_cell="A2",
                        header_fill="1F4E78", header_font_color="FFFFFF",
                        max_col=33)
    _destacar_totais(ws_deb, [resumo_row, subtotal_row, tg_row, check_row],
                      cor_fundo="D9E1F2", min_col=11, max_col=23)

    log("\n[8] Salvando...")
    wb.save(saida)
    log(f"\nCONCLUIDO! Arquivo: {os.path.basename(saida)}")

    # Calcular o % proporcional AQUI em Python (mesma conta da fórmula
    # FAT BRUTO / VEIC USADO / DEMAIS) e devolver como número pronto —
    # não depender de reabrir o .xlsx num motor de planilha depois, já
    # que fórmulas escritas via openpyxl não têm valor em cache até
    # serem calculadas por Excel/LibreOffice pelo menos uma vez.
    fat_bruto_total = 0.0
    valor_veiculo_usado = 0.0
    for _, row in combined.iterrows():
        fat = row.iloc[17]  # col S = Fat. Bruto (índice pós-inserção de Operação/Tipo Saida)
        try: fat_v = float(fat) if pd.notna(fat) else 0.0
        except: fat_v = 0.0
        fat_bruto_total += fat_v

        tipo_h = str(row.iloc[7]) if pd.notna(row.iloc[7]) else ''
        if tipo_h in ('Garantia', 'Ativo Imobilizado', 'Devolução', 'Venda Direta'):
            tipo_y = 'Zero'
        elif tipo_h in ('Comissão', 'Oficina', 'Serviço - A Classificar', 'Combustível', 'Material de Uso e Consumo'):
            tipo_y = 'Normal NC'
        elif tipo_h == 'Veículo Novo':
            tipo_y = 'Monofasico'
        elif tipo_h == 'Veículo Usado':
            tipo_y = 'Normal C'
        else:
            ncm_str2 = limpar_ncm(row.iloc[4])
            if ncm_str2 in existing_ncms:
                tipo_y = existing_ncms[ncm_str2][1] or 'Normal NC'
            elif ncm_str2 in missing_ncms:
                tipo_y = missing_ncms[ncm_str2][0]
            else:
                tipo_y = 'Normal NC'
        chave = chave_resumo(tipo_h, tipo_y)
        if chave == 'VEICULO USADO - NORMAL C':
            valor_veiculo_usado += fat_v

    valor_demais = fat_bruto_total - valor_veiculo_usado
    pct_proporcional = (valor_demais / fat_bruto_total) if fat_bruto_total else 0.0
    log(f"\n   Proporcional calculado em Python: FAT BRUTO={fat_bruto_total:,.2f} | "
        f"VEIC USADO={valor_veiculo_usado:,.2f} | DEMAIS={valor_demais:,.2f} | % = {pct_proporcional:.4%}")

    return {
        "arquivo": saida,
        "fat_bruto_total": fat_bruto_total,
        "valor_veiculo_usado": valor_veiculo_usado,
        "valor_demais": valor_demais,
        "pct_proporcional": pct_proporcional,
    }
